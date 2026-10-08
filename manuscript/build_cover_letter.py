# -*- coding: utf-8 -*-
"""Cover letter for Reaction Kinetics, Mechanisms and Catalysis (Springer), original article.   python build_cover_letter.py -> out/Cover_Letter_RKMC.docx"""
import json
import os
import sys

from docx.shared import Pt

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
import docx_helpers as H  # noqa: E402

ZEN = json.load(open(os.path.join("C:" + os.sep, "YouTube", "_gac_zenodo_state.json")))
SW, PP = (ZEN.get("software_1.3.0") or ZEN.get("software_1.2.0") or ZEN["software"])["doi"], (ZEN.get("publication_v4") or ZEN.get("publication_v3") or ZEN["publication"])["doi"]
R = json.load(open(os.path.join(HERE, "..", "results.json"), encoding="utf-8"))
REPO = "https://github.com/sandlerleon/gated-autocatalysis-criticality"
TITLE = "Closed-Form Storage Stability, Induction Delay, and Thermal-Feedback Criticality of Melt-Gated Autocatalytic Cure"
dr = [r for r in R["design_rule"] if r["sigma_star_ode"] and r["sigma_star_closed"] and r["sigma_star_closed"] < 40]
dev = max(abs(r["sigma_star_closed"] / r["sigma_star_ode"] - 1) for r in dr)
case20 = [c for c in R["flow"]["cases"] if c["psi"] == 20.0][0]
win220 = [c for c in R["arrhenius_flow"]["cases"] if c["dtad"] == 220.0][0]["analysis"]["window"]["width_K"]
doc = H.new_document(size=11, line=1.15)


def para(text, bold=False, after=6):
    p = doc.add_paragraph()
    if bold:
        p.add_run(text).bold = True
    else:
        H.add_rich(p, text)
    p.paragraph_format.space_after = Pt(after)
    return p


def bullet(text):
    p = doc.add_paragraph(style="List Bullet")
    H.add_rich(p, text)
    p.paragraph_format.space_after = Pt(3)


for line in ("Leon Sandler", "Independent researcher, Northbrook, Illinois, USA", "sandler.leon@gmail.com", "ORCID: https://orcid.org/0009-0007-4584-808X"):
    para(line, after=0)
para("")
para("7 October 2026")
para("The Editors\nReaction Kinetics, Mechanisms and Catalysis", after=10)
para("Submission of an original article: “%s”" % TITLE, bold=True, after=10)
para("Dear Editors,")
para("I submit the enclosed theoretical paper for consideration in Reaction Kinetics, Mechanisms and Catalysis. "
     "It concerns the kinetics of a cure that is triggered by the melting of an encapsulated catalyst, with a distribution of capsule melting temperatures, autocatalytic (Kamal–Sourour) rate laws and exothermic feedback, "
     "and the behavior of that reaction in a batch and in a continuous-flow reactor. The paper is purely theoretical: no experimental data are used and all parameters are illustrative.")
para("What is new, and what is not", bold=True, after=3)
para("Several ingredients are standard, and the paper says so: separable cure kinetics with a generalized time, the Semenov and Frank-Kamenetskii fold and its extension to reactant consumption, and the "
     "multiplicity and stability analysis of continuous stirred reactors. The new content is their assembly for a melt-gated autocatalytic cure, with results of three different kinds that the manuscript "
     "keeps apart:", after=3)
bullet("*Exact results.* At constant temperature the gated cure is the isothermal cure on the clock of integrated catalyst availability; for a fully activated population the delay is "
       "τ_{rel}[1 − exp(−t/τ_{rel})], which depends on conversion and approaches the release time constant only at late conversion. The isothermal cure time is exact by quadrature. The steady states of the flow reactor and the Jacobian are exact for the stated model.")
bullet("*Closed forms with stated approximations.* The storage conversion of a Gaussian capsule population, (a/b)[exp(b f t_{eff}) − 1] with f = Φ(−Δ/σ_{eff}), and the bound on the melting-temperature spread that "
       "follows from it (agreement with the full population equations within %.0f%% over the tested grid). The frozen-conversion critical number (1+n)^{1+n}/[e n^{n}(1+ε)^{1+n}] of the thermal balance, with an Arrhenius correction that agrees with integrations "
       "to a few percent over the tested range." % (100 * dev))
bullet("*Numerical evidence.* Operational ignition thresholds from integrations; the window of two stable states of the flow reactor, whose upper branch can lose stability at a Hopf-type point before its fold "
       "(ψ = 20: the stable window starts at ln D = %.2f rather than at the fold; sweep loops through that point depend on noise and sweep rate and are not shown to converge to the attractor loop area %.2f; an Arrhenius temperature-path simulation gives a window of %.1f K for the illustrative parameters). The manuscript states where each result has been tested and where it has not." % (
           case20["window"]["lnD_ext"], case20["attractor_area"], win220))
para("Fit to the journal", bold=True, after=3)
para("The paper links reaction kinetics to reactor behavior: it shows how quantities that a calorimeter measures (rate constants, the width of the melting endotherm, a thermal time constant) enter the storage, delay and "
     "runaway behavior of a latent catalytic system through a small number of closed-form relations, with their domains of validity and the signatures that would test them. Whether these model reductions "
     "constitute a sufficient advance for the journal is for the Editors to judge; I have tried to make the claims no stronger than the checks that support them.")
para("Relation to related work by the author (disclosure)", bold=True, after=3)
para("The model equations, the illustrative parameter set, the capsule-population picture, the logistic activation and the reference integrator are those of an earlier numerical study of mine on coupled thermal activation and "
     "autocatalytic cure. That study is available as a ChemRxiv preprint (10.26434/chemrxiv.15008366; Zenodo 10.5281/zenodo.22073390) and has been submitted for publication in another journal; a later version (Zenodo 10.5281/zenodo.23202125) added numerical "
     "storage-stability and design-rule tables. Both are cited in the manuscript (Section 1.1) and are not claimed as new here: the coupled kinetic model is not presented as a contribution of this paper, no figure of the earlier "
     "study is reproduced, its numerical comparisons with first-order controls are not used, and its design-rule tables are used only as an independent test of the closed form of Proposition 3. "
     "The present paper asks a different question, namely which statements about storage, delay, criticality and reactor stability can be derived analytically, and Table 1 of the manuscript lists the contributions that are specific to it. "
     "I disclose the related submission here so that the Editors can judge the overlap, and I will gladly provide the earlier manuscript on request.")
para("Data, code and disclosures", bold=True, after=3)
para("All code, tests, raw result tables and figure scripts are public at %s (release v1.3.0) and archived at https://doi.org/%s; the manuscript is available as a preprint at https://doi.org/%s. "
     "The manuscript describes in Section 2.3 and in the Statements and Declarations the use of Claude Sonnet 5.5 (Anthropic; model identifier as reported by the Claude Code environment of the Claude desktop application) for derivations, code, figures and drafting, and states that every number "
     "is produced by the released scripts and that I am responsible for the content. I am the sole author, have no competing interests and received no funding. The manuscript is not under consideration elsewhere." % (REPO, SW, PP))
para("Thank you for considering the paper.")
para("Yours sincerely,", after=18)
para("Leon Sandler")
out = os.path.join(HERE, "out", "Cover_Letter_RKMC.docx")
doc.save(out)
print("saved", out)
